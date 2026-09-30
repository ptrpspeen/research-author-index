#!/usr/bin/env python3
"""Validate researched article records and export a two-sheet Excel workbook."""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
from datetime import date
from pathlib import Path
from typing import Optional
from urllib.parse import parse_qs, urlparse


COMMON = [
    ("researcher", "ผู้วิจัย", 25),
    ("affiliation", "หน่วยงานผู้วิจัย", 27),
    ("researcher_orcid", "ORCID ผู้วิจัย", 24),
    ("researcher_orcid_url", "ลิงก์ ORCID", 48),
    ("title", "ชื่อบทความ", 55),
    ("publication_year", "ปีตีพิมพ์ (ค.ศ.)", 19),
    ("document_type", "ประเภทเอกสาร", 21),
    ("authors", "ผู้แต่งทั้งหมด", 55),
    ("authors_complete", "รายชื่อผู้แต่งครบ", 20),
    ("journal", "วารสาร/แหล่งตีพิมพ์", 42),
    ("issn", "ISSN", 17),
    ("doi", "DOI", 28),
]
SCOPUS_EXTRA = [
    ("scopus_eid", "Scopus EID", 26),
    ("scopus_author_id", "Scopus Author ID", 22),
    ("author_h_index", "h-index ผู้วิจัย", 19),
    ("h_index_as_of", "วันที่ตรวจ h-index", 21),
    ("citations", "อ้างอิงบทความ", 19),
    ("citations_as_of", "วันที่ตรวจอ้างอิง", 21),
    ("quartile", "Q (CiteScore)", 18),
    ("citescore_percentile", "CiteScore percentile", 22),
    ("quartile_year", "ปีของ Q", 13),
    ("quartile_subject", "สาขาของ Q", 33),
    ("quartile_source_url", "หลักฐาน Q", 55),
    ("author_profile_url", "ลิงก์โปรไฟล์ผู้วิจัย", 55),
]
TCI_EXTRA = [
    ("volume", "ปีที่/เล่มที่", 16),
    ("issue", "ฉบับที่", 14),
    ("issue_title", "ชื่อฉบับ", 43),
    ("issue_url", "ลิงก์ฉบับ", 58),
    ("tier", "Tier ในปีตีพิมพ์", 20),
    ("tier_period", "ช่วง/วันที่ของ Tier", 27),
    ("tier_source_url", "หลักฐาน Tier", 55),
    ("tier_evidence_basis", "ลักษณะหลักฐาน Tier", 24),
    ("current_tier", "Tier ปัจจุบัน ณ วันที่ตรวจ", 27),
    ("current_tier_checked_at", "วันที่ตรวจ Tier ปัจจุบัน", 24),
    ("current_tier_source_url", "หลักฐาน Tier ปัจจุบัน", 55),
]
TAIL = [
    ("article_url", "ลิงก์บทความ", 65),
    ("retrieved_at", "วันที่สืบค้น", 17),
    ("verification_status", "สถานะตรวจสอบ", 20),
    ("notes", "หมายเหตุ", 48),
]
URL_FIELDS = {
    "article_url", "issue_url", "author_profile_url", "quartile_source_url", "tier_source_url", "current_tier_source_url", "researcher_orcid_url"
}
DATE_FIELDS = {"retrieved_at", "h_index_as_of", "citations_as_of", "current_tier_checked_at"}


def fail(where: str, message: str) -> None:
    raise ValueError(f"{where}: {message}")


def require_text(value: object, where: str) -> str:
    if not isinstance(value, str) or not value.strip():
        fail(where, "ต้องเป็นข้อความที่ไม่ว่าง")
    return value.strip()


def optional_text(value: object, where: str) -> str:
    if value is None or value == "":
        return ""
    return require_text(value, where)


def check_date(value: object, where: str) -> str:
    text = require_text(value, where)
    try:
        parsed = date.fromisoformat(text)
    except ValueError:
        fail(where, "ต้องเป็นวันที่ YYYY-MM-DD")
    if parsed.isoformat() != text:
        fail(where, "ต้องเป็นวันที่ YYYY-MM-DD")
    return text


def check_url(value: object, where: str, required: bool = False) -> str:
    text = require_text(value, where) if required else optional_text(value, where)
    if not text:
        return ""
    parsed = urlparse(text)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username:
        fail(where, "ต้องเป็นลิงก์ HTTP(S) ที่มีโดเมน และไม่มีชื่อผู้ใช้")
    if len(text) > 2000:
        fail(where, "ลิงก์ยาวเกิน 2,000 อักขระ")
    return text


def require_domain(url: str, domain: str, where: str) -> None:
    if url:
        host = urlparse(url).hostname or ""
        if host != domain and not host.endswith("." + domain):
            fail(where, f"หลักฐานต้องอยู่บนโดเมน {domain}")


def validate_tier_source(url: str, article_url: str, where: str) -> None:
    if not url:
        return
    parsed = urlparse(url)
    host = parsed.hostname or ""
    if host == "tci-thailand.org" or host.endswith(".tci-thailand.org"):
        return
    article_match = re.match(r"^/index\.php/([^/]+)/article/view/", urlparse(article_url).path)
    slug = article_match.group(1).casefold() if article_match else ""
    article_host = urlparse(article_url).hostname
    prefix = f"/index.php/{slug}"
    if parsed.scheme == "https" and host == article_host and (parsed.path.casefold() == prefix or parsed.path.casefold().startswith(prefix + "/")):
        return
    if host in {"tci-thaijo.org", "www.tci-thaijo.org"}:
        if parsed.path == "/public/tinfo.php" and parse_qs(parsed.query).get("path", [""])[0].casefold() == slug:
            return
        if parsed.path.casefold() in {f"/journals/{slug}", f"/en/journals/{slug}"}:
            return
    fail(where, "หลักฐาน Tier ต้องเป็น TCI Journal List หรือภาพตรา/หน้า ThaiJO ของวารสารเดียวกับบทความ")


def certified_year_range(period: str) -> tuple[int, int] | None:
    match = re.search(r"(?<!\d)((?:20|25)\d{2})(?:-\d{2}-\d{2})?\s*(?:-|–|—|\.\.|ถึง|to)\s*((?:20|25)\d{2})(?!\d)", period, re.IGNORECASE)
    if not match:
        return None
    start, end = (int(value) for value in match.groups())
    if start >= 2400 and end >= 2400:
        start, end = start - 543, end - 543
    return start, end


def check_int(value: object, where: str, minimum: int = 0) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value < minimum:
        fail(where, f"ต้องเป็นจำนวนเต็มตั้งแต่ {minimum} ขึ้นไป")
    return value


def optional_int(value: object, where: str) -> Optional[int]:
    if value is None or value == "":
        return None
    return check_int(value, where)


def check_year(value: object, where: str) -> int:
    year = check_int(value, where, 1900)
    if year > date.today().year + 1:
        fail(where, "ปีเกินปีปัจจุบัน+1; แปลง พ.ศ. เป็น ค.ศ. ก่อน")
    return year


def check_orcid(value: object, where: str) -> str:
    identifier = optional_text(value, where)
    if not identifier:
        return ""
    if identifier.startswith("https://"):
        parsed = urlparse(identifier)
        if (parsed.netloc != "orcid.org" or parsed.query or parsed.fragment
                or not re.fullmatch(r"/\d{4}-\d{4}-\d{4}-\d{3}[\dX]/?", parsed.path)):
            fail(where, "ORCID URL must use https://orcid.org/ followed by one identifier")
        identifier = parsed.path.strip("/")
    if not re.fullmatch(r"\d{4}-\d{4}-\d{4}-\d{3}[\dX]", identifier):
        fail(where, "รูปแบบ ORCID ไม่ถูกต้อง")
    compact = identifier.replace("-", "")
    total = 0
    for digit in compact[:15]:
        total = (total + int(digit)) * 2
    check = (12 - total % 11) % 11
    expected = "X" if check == 10 else str(check)
    if compact[-1] != expected:
        fail(where, "เลขตรวจสอบ ORCID ไม่ถูกต้อง")
    return identifier


def quartile_from_percentile(percentile: int) -> str:
    if percentile >= 75:
        return "Q1"
    if percentile >= 50:
        return "Q2"
    if percentile >= 25:
        return "Q3"
    return "Q4"


def is_journal_document(row: dict) -> bool:
    return row.get("document_type", "").casefold() in {"article", "article in press", "review"}


def scopus_link_eid(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or parsed.hostname not in {"scopus.com", "www.scopus.com"}:
        return ""
    match = re.fullmatch(r"/pages/publications/(\d+)/?", parsed.path)
    if match:
        return "2-s2.0-" + match.group(1)
    if parsed.path in {"/record/display.uri", "/inward/record.uri"}:
        values = parse_qs(parsed.query).get("eid", [])
        if len(values) == 1 and re.fullmatch(r"2-s2\.0-\d+", values[0]):
            return values[0]
    return ""


def validate_record(record: object, database: str, index: int, retrieved_year: int | None = None) -> dict:
    where = f"{database}[{index}]"
    if not isinstance(record, dict):
        fail(where, "ต้องเป็น object")
    row = dict(record)
    row["title"] = require_text(row.get("title"), f"{where}.title")
    row["publication_year"] = check_year(row.get("publication_year"), f"{where}.publication_year")
    row["journal"] = require_text(row.get("journal"), f"{where}.journal")
    authors = row.get("authors")
    if not isinstance(authors, list) or not authors:
        fail(f"{where}.authors", "ต้องเป็นรายชื่อผู้แต่งอย่างน้อยหนึ่งคน")
    row["authors"] = [require_text(name, f"{where}.authors[{i}]") for i, name in enumerate(authors)]
    if type(row.get("authors_complete")) is not bool:
        fail(f"{where}.authors_complete", "ต้องเป็น true หรือ false")
    row["article_url"] = check_url(row.get("article_url"), f"{where}.article_url", True)
    if database == "tci":
        require_domain(row["article_url"], "tci-thaijo.org", f"{where}.article_url")
        if "/article/view/" not in urlparse(row["article_url"]).path:
            fail(f"{where}.article_url", "ต้องเป็นหน้าบทความ ThaiJO โดยตรง")
    else:
        require_domain(row["article_url"], "scopus.com", f"{where}.article_url")
    for field in ("issn", "doi", "document_type", "notes"):
        row[field] = optional_text(row.get(field), f"{where}.{field}")
    status = row.get("verification_status", "confirmed")
    if status not in {"confirmed", "needs_review"}:
        fail(f"{where}.verification_status", "ใช้ confirmed หรือ needs_review")
    if not row["authors_complete"] and status != "needs_review":
        fail(f"{where}.verification_status", "ผู้แต่งไม่ครบต้องใช้ needs_review")
    row["verification_status"] = status

    if database == "scopus":
        for field in ("scopus_eid", "scopus_author_id", "quartile_subject"):
            row[field] = optional_text(row.get(field), f"{where}.{field}")
        link_eid = scopus_link_eid(row["article_url"])
        if not link_eid:
            fail(f"{where}.article_url", "ต้องเป็นหน้าบทความ Scopus โดยตรง")
        if row["scopus_eid"]:
            match = re.fullmatch(r"2-s2\.0-(\d+)", row["scopus_eid"])
            if not match:
                fail(f"{where}.scopus_eid", "รูปแบบ EID ไม่ถูกต้อง")
            if link_eid != row["scopus_eid"]:
                fail(f"{where}.article_url", "ลิงก์ไม่ตรงกับ EID")
        for field in ("author_profile_url", "quartile_source_url"):
            row[field] = check_url(row.get(field), f"{where}.{field}")
            require_domain(row[field], "scopus.com", f"{where}.{field}")
        if row["scopus_author_id"] and row["author_profile_url"]:
            profile = urlparse(row["author_profile_url"])
            ids = parse_qs(profile.query).get("authorId", [])
            modern = re.fullmatch(r"/pages/authors/(\d+)/?", profile.path)
            if not ((profile.path == "/authid/detail.uri" and ids == [row["scopus_author_id"]])
                    or (modern and modern.group(1) == row["scopus_author_id"])):
                fail(f"{where}.author_profile_url", "ลิงก์โปรไฟล์ไม่ตรงกับ Scopus Author ID")
        for field in ("author_h_index", "citations"):
            row[field] = optional_int(row.get(field), f"{where}.{field}")
        for field in ("h_index_as_of", "citations_as_of"):
            if row.get(field):
                row[field] = check_date(row[field], f"{where}.{field}")
        if row["author_h_index"] is not None and (
            not row.get("h_index_as_of") or not row["author_profile_url"]
        ):
            fail(where, "h-index ต้องมีวันที่ตรวจและ URL โปรไฟล์")
        if row["citations"] is not None and not row.get("citations_as_of"):
            fail(where, "จำนวนอ้างอิงต้องมีวันที่ตรวจ")
        percentile = optional_int(row.get("citescore_percentile"), f"{where}.citescore_percentile")
        if percentile is not None and percentile > 100:
            fail(f"{where}.citescore_percentile", "ต้องไม่เกิน 100")
        row["citescore_percentile"] = percentile
        quartile = row.get("quartile") or ""
        if quartile and quartile not in {"Q1", "Q2", "Q3", "Q4"}:
            fail(f"{where}.quartile", "ใช้ Q1, Q2, Q3 หรือ Q4")
        if percentile is not None:
            calculated = quartile_from_percentile(percentile)
            if quartile and quartile != calculated:
                fail(f"{where}.quartile", f"ไม่ตรงกับ CiteScore percentile {percentile} ({calculated})")
            quartile = calculated
        row["quartile"] = quartile
        if quartile:
            if not is_journal_document(row):
                fail(f"{where}.quartile", "Q วารสารใช้กับ Article/Review; ไม่เติมให้ Conference Paper")
            row["quartile_year"] = check_year(row.get("quartile_year"), f"{where}.quartile_year")
            prior_current = (row["publication_year"] == (retrieved_year or date.today().year)
                             and row["quartile_year"] == row["publication_year"] - 1 and bool(row["notes"]))
            if row["quartile_year"] != row["publication_year"] and not prior_current:
                fail(f"{where}.quartile_year", "Q ต้องตรงปีบทความ; ปีปัจจุบันใช้ปีก่อนหน้าได้พร้อมหมายเหตุหลักฐานปีปิดรอบ")
            if not row["quartile_subject"] or not row["quartile_source_url"]:
                fail(where, "Q ต้องมีสาขาและ URL หลักฐาน")
    else:
        for field in ("volume", "issue", "issue_title"):
            row[field] = optional_text(row.get(field), f"{where}.{field}")
        row["issue_url"] = check_url(row.get("issue_url"), f"{where}.issue_url")
        require_domain(row["issue_url"], "tci-thaijo.org", f"{where}.issue_url")
        if row["issue_url"]:
            article_parsed = urlparse(row["article_url"])
            issue_parsed = urlparse(row["issue_url"])
            article_match = re.match(r"^/index\.php/([^/]+)/article/view/", article_parsed.path)
            slug = article_match.group(1) if article_match else ""
            if issue_parsed.hostname != article_parsed.hostname or not issue_parsed.path.startswith(f"/index.php/{slug}/issue/view/"):
                fail(f"{where}.issue_url", "ต้องเป็นหน้าฉบับของวารสารเดียวกับบทความ")
        row["tier_source_url"] = check_url(row.get("tier_source_url"), f"{where}.tier_source_url")
        validate_tier_source(row["tier_source_url"], row["article_url"], f"{where}.tier_source_url")
        row["tier_period"] = optional_text(row.get("tier_period"), f"{where}.tier_period")
        basis = optional_text(row.get("tier_evidence_basis"), f"{where}.tier_evidence_basis")
        if basis not in {"", "explicit_period", "publisher_since"}:
            fail(where, "unknown Tier evidence basis")
        row["tier_evidence_basis"] = basis
        current = row.get("current_tier") or ""
        if current and current not in {"Tier 1", "Tier 2", "Tier 3"}:
            fail(where, "invalid current Tier")
        row["current_tier"] = current
        row["current_tier_source_url"] = check_url(row.get("current_tier_source_url"), f"{where}.current_tier_source_url")
        validate_tier_source(row["current_tier_source_url"], row["article_url"], f"{where}.current_tier_source_url")
        row["current_tier_checked_at"] = optional_text(row.get("current_tier_checked_at"), f"{where}.current_tier_checked_at")
        if current:
            check_date(row["current_tier_checked_at"], f"{where}.current_tier_checked_at")
            if not row["current_tier_source_url"] or row["current_tier_checked_at"] > date.today().isoformat():
                fail(where, "current Tier needs source and non-future observation date")
        tier = row.get("tier") or ""
        if tier and tier not in {"Tier 1", "Tier 2", "Tier 3"}:
            fail(f"{where}.tier", "ใช้ Tier 1, Tier 2 หรือ Tier 3")
        row["tier"] = tier
        if tier and (not row["tier_period"] or not row["tier_source_url"]):
            fail(where, "Tier ต้องมีช่วง/วันที่และ URL หลักฐาน")
        years = certified_year_range(row["tier_period"])
        if tier and not years:
            fail(f"{where}.tier_period", "Tier ต้องมีช่วงปีรับรอง; วันที่ตรวจอย่างเดียวไม่ใช่หลักฐานย้อนหลัง")
        if tier and basis == "publisher_since":
            observed = re.search(r"observed (\d{4}-\d{2}-\d{2}); no certified expiry", row["tier_period"])
            if not observed or not years:
                fail(where, "publisher_since needs its observation date and no certified expiry label")
            check_date(observed[1], f"{where}.tier_period")
            if years[1] != int(observed[1][:4]) or observed[1] > date.today().isoformat():
                fail(where, "publisher_since must stop at its observation year")
        if tier and years and not years[0] <= row["publication_year"] <= years[1]:
            fail(f"{where}.tier_period", "ปีบทความอยู่นอกช่วงรับรอง Tier ที่ระบุ")
    return row


def duplicate_keys(row: dict, database: str) -> set[str]:
    # Repeated editorial titles within a year can identify different issues.
    # Use stable identifiers for hard rejection; title/year is a review signal.
    keys: set[str] = set()
    if database == "tci" and row.get("doi"):
        keys.add("doi:" + normalized_doi(row["doi"]))
    if database == "scopus" and row.get("scopus_eid"):
        keys.add("eid:" + row["scopus_eid"].lower())
    keys.add("url:" + row["article_url"].rstrip("/").lower())
    return keys


def normalized_doi(value: str) -> str:
    doi = value.strip().casefold()
    return re.sub(r"^(?:https?://(?:dx\.)?doi\.org/|doi:\s*)", "", doi).rstrip("/ ")


def scopus_doi_conflicts(rows: list[dict]) -> dict[str, list[dict]]:
    by_doi: dict[str, list[dict]] = {}
    for row in rows:
        if row.get("doi"):
            by_doi.setdefault(normalized_doi(row["doi"]), []).append(row)
    return {doi: group for doi, group in by_doi.items()
            if len({row.get("scopus_eid") for row in group if row.get("scopus_eid")}) > 1}


def annotate_scopus_doi_conflicts(rows: list[dict]) -> list[dict]:
    for doi, group in scopus_doi_conflicts(rows).items():
        eids = ", ".join(sorted(row["scopus_eid"] for row in group if row.get("scopus_eid")))
        note = (f"DOI conflict: shared by distinct Scopus EIDs {eids}; verify whether these records "
                "represent one work before counting unique works.")
        for row in group:
            row["verification_status"] = "needs_review"
            current = row.get("notes", "").strip()
            if note not in current:
                row["notes"] = "; ".join(part for part in (current, note) if part)
    return rows


def validate_payload(payload: object) -> dict:
    if not isinstance(payload, dict):
        fail("JSON", "ต้องเป็น object")
    researcher = payload.get("researcher")
    if not isinstance(researcher, dict):
        fail("researcher", "ต้องเป็น object")
    name = require_text(researcher.get("name"), "researcher.name")
    affiliation = optional_text(researcher.get("affiliation"), "researcher.affiliation")
    researcher_orcid = check_orcid(researcher.get("orcid"), "researcher.orcid")
    retrieved_at = check_date(payload.get("retrieved_at"), "retrieved_at")
    result = {"name": name, "affiliation": affiliation, "researcher_orcid": researcher_orcid,
              "researcher_orcid_url": f"https://orcid.org/{researcher_orcid}" if researcher_orcid else "",
              "retrieved_at": retrieved_at}
    for database in ("scopus", "tci"):
        records = payload.get(database)
        if not isinstance(records, list):
            fail(database, "ต้องเป็น array (ถ้าไม่มีผลงานให้ใช้ [])")
        result[database] = []
        seen: set[str] = set()
        for index, item in enumerate(records):
            row = validate_record(item, database, index, int(retrieved_at[:4]))
            common = duplicate_keys(row, database) & seen
            if common:
                fail(f"{database}[{index}]", "พบรายการซ้ำในชีต: " + sorted(common)[0])
            seen.update(duplicate_keys(row, database))
            result[database].append(row)
    annotate_scopus_doi_conflicts(result["scopus"])
    if not result["scopus"] and not result["tci"]:
        fail("JSON", "ไม่มีบทความในทั้งสองฐาน; อย่าสร้างรายงานเปล่า")
    audit = payload.get("audit", {})
    if not isinstance(audit, dict):
        fail("audit", "ต้องเป็น object")
    result["audit"] = audit
    searches = audit.get("searches", [])
    exclusions = audit.get("exclusions", [])
    if not isinstance(searches, list) or not isinstance(exclusions, list):
        fail("audit", "searches และ exclusions ต้องเป็น array")
    for index, entry in enumerate(searches):
        if not isinstance(entry, dict) or entry.get("database") not in {"scopus", "tci"}:
            fail(f"audit.searches[{index}]", "ต้องระบุ database เป็น scopus หรือ tci")
        require_text(entry.get("query"), f"audit.searches[{index}].query")
        check_int(entry.get("results"), f"audit.searches[{index}].results")
        check_date(entry.get("checked_at"), f"audit.searches[{index}].checked_at")
    for index, entry in enumerate(exclusions):
        if not isinstance(entry, dict) or entry.get("database") not in {"scopus", "tci"}:
            fail(f"audit.exclusions[{index}]", "ต้องระบุ database เป็น scopus หรือ tci")
        require_text(entry.get("reason"), f"audit.exclusions[{index}].reason")
        require_text(entry.get("identifier"), f"audit.exclusions[{index}].identifier")
    for database in ("scopus", "tci"):
        expected = audit.get(f"{database}_expected_count")
        excluded = audit.get(f"{database}_excluded_count", 0)
        if expected is not None:
            expected = check_int(expected, f"audit.{database}_expected_count")
            excluded = check_int(excluded, f"audit.{database}_excluded_count")
            if excluded and sum(item.get("database") == database for item in exclusions) != excluded:
                fail("audit", f"{database}: จำนวน exclusions ไม่ตรงกับเหตุผลที่บันทึก")
            if len(result[database]) != expected - excluded:
                fail("audit", f"{database}: {len(result[database])} แถว ไม่ตรงกับผลค้น {expected} ลบรายการที่ตัดออก {excluded}")
    return result


def preflight_summary(payload: dict) -> dict:
    warnings = []
    doi_conflicts = scopus_doi_conflicts(payload["scopus"])
    for doi, group in doi_conflicts.items():
        eids = ", ".join(sorted(row["scopus_eid"] for row in group if row.get("scopus_eid")))
        warnings.append(f"Scopus DOI {doi} is shared by distinct EIDs {eids}; observed records are not a unique-work count until reviewed")
    for database in ("scopus", "tci"):
        titles: dict[tuple, int] = {}
        for number, row in enumerate(payload[database], 1):
            label = f"{database}[{number}] {row['title']}"
            title_key = (re.sub(r"\s+", " ", row["title"]).casefold(), row["publication_year"], row["journal"])
            if title_key in titles:
                warnings.append(f"{label}: ชื่อ/ปี/วารสารตรงกับแถว {titles[title_key]} แต่รหัสต่างกัน; ตรวจฉบับก่อนพิจารณาซ้ำ")
            else:
                titles[title_key] = number
            if not row["document_type"]:
                warnings.append(f"{label}: ยังไม่ยืนยันประเภทเอกสาร")
            if not row["authors_complete"]:
                warnings.append(f"{label}: รายชื่อผู้แต่งยังไม่ครบ")
            if row["verification_status"] == "needs_review":
                warnings.append(f"{label}: ยังต้องตรวจทาน")
            if database == "scopus" and is_journal_document(row) and not row["quartile"]:
                warnings.append(f"{label}: ยังไม่มี Q พร้อมหลักฐานปี/สาขา")
            if database == "tci" and not row["tier"]:
                warnings.append(f"{label}: ยังไม่มี Tier พร้อมหลักฐานช่วงเวลา")
    audit = payload["audit"]
    q_eligible = [row for row in payload["scopus"] if is_journal_document(row)]
    q_filled = sum(bool(row.get("quartile")) for row in q_eligible)
    tier_filled = sum(bool(row.get("tier")) for row in payload["tci"])
    return {"scopus_rows": len(payload["scopus"]), "tci_rows": len(payload["tci"]),
            "scopus_doi_conflict_groups": len(doi_conflicts),
            "scopus_q_eligible_rows": len(q_eligible),
            "scopus_q_evidenced_rows": q_filled,
            "scopus_q_missing_eligible_rows": len(q_eligible) - q_filled,
            "scopus_q_not_applicable_rows": len(payload["scopus"]) - len(q_eligible),
            "tci_tier_evidenced_rows": tier_filled,
            "tci_current_tier_evidenced_rows": sum(bool(r.get("current_tier")) for r in payload["tci"]),
            "tci_tier_publisher_since_rows": sum(r.get("tier_evidence_basis") == "publisher_since" for r in payload["tci"]),
            "tci_tier_missing_rows": len(payload["tci"]) - tier_filled,
            "searches_logged": len(audit.get("searches", [])),
            "exclusions_logged": len(audit.get("exclusions", [])),
            "warnings": warnings, "warning_count": len(warnings)}


def safe_cell(cell, value: object) -> None:
    if value is None:
        return
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Strings such as '=...', '+...' remain text, never formulas.


def make_sheet(workbook, name: str, columns: list[tuple], rows: list[dict], meta: dict) -> None:
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    ws = workbook.create_sheet(name)
    for col_num, (field, label, width) in enumerate(columns, 1):
        header = ws.cell(1, col_num, label)
        header.font = Font(color="FFFFFF", bold=True)
        header.fill = PatternFill("solid", fgColor="164A70")
        header.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_num)].width = width
        if field == "author_h_index":
            header.comment = Comment("ค่า h-index ของผู้วิจัยใน Scopus ณ วันที่ตรวจ ไม่ใช่ค่าของบทความ", "Research Author Index")
        elif field in {"quartile", "tier", "current_tier"}:
            header.comment = Comment("ระดับของวารสาร โปรดดูปี/ช่วงเวลาและ URL หลักฐานในแถวเดียวกัน", "Research Author Index")
    ws.row_dimensions[1].height = 32
    for row_num, row in enumerate(rows, 2):
        context = dict(row, researcher=meta["name"], affiliation=meta["affiliation"], retrieved_at=meta["retrieved_at"],
                       researcher_orcid=meta.get("researcher_orcid", ""),
                       researcher_orcid_url=meta.get("researcher_orcid_url", ""))
        context["authors"] = "; ".join(row["authors"])
        context["authors_complete"] = "ครบ" if row["authors_complete"] else "ไม่ครบ"
        for col_num, (field, _, _) in enumerate(columns, 1):
            cell = ws.cell(row_num, col_num)
            value = context.get(field)
            safe_cell(cell, value)
            cell.alignment = Alignment(vertical="top", wrap_text=field in {"title", "authors", "notes"})
            if field in URL_FIELDS and value:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
        ws.row_dimensions[row_num].height = 42
    ws.freeze_panes = "D2"
    last_col = get_column_letter(len(columns))
    last_row = max(1, len(rows) + 1)
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    if rows:
        table = Table(displayName=f"{name}Papers", ref=f"A1:{last_col}{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    ws.sheet_view.showGridLines = False


def export(payload: dict, output: Path, allow_partial: bool = False) -> None:
    audit = payload.get("audit", {})
    verified_zero = (audit.get("scopus_expected_count") == 0 and
                     any(s.get("database") == "scopus" and s.get("results") == 0
                         for s in audit.get("searches", [])))
    partial = not payload["scopus"] and not verified_zero
    if partial and not allow_partial:
        fail("Scopus", "ยังไม่ยืนยันว่าไม่มีผลงาน: เก็บ draft ไว้; --allow-partial ใช้เฉพาะเมื่อผู้ใช้ยอมรับรายงานที่ยังไม่ครบ")
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("ต้องติดตั้ง openpyxl: python3 -m pip install 'openpyxl>=3.1,<4' (Windows: py -3 -m pip install 'openpyxl>=3.1,<4')") from exc
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = f"ผลงานวิจัย: {payload['name']}"
    workbook.properties.description = f"Scopus และ ThaiJO/TCI; สืบค้น {payload['retrieved_at']}"
    doi_conflicts = scopus_doi_conflicts(payload["scopus"])
    if doi_conflicts:
        workbook.properties.description += (
            f"; Scopus มี {len(doi_conflicts)} กลุ่ม DOI conflict; จำนวนแถวเป็นจำนวน records ไม่ใช่จำนวน unique works"
        )
    if partial:
        workbook.properties.title = "รายงานบางส่วน: " + workbook.properties.title
        workbook.properties.description += "; รายงานบางส่วน: Scopus ยังไม่ครบ ไม่ใช่ยืนยันว่าไม่มีผลงาน"
    make_sheet(workbook, "Scopus", COMMON + SCOPUS_EXTRA + TAIL, payload["scopus"], payload)
    make_sheet(workbook, "TCI", COMMON + TCI_EXTRA + TAIL, payload["tci"], payload)
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.suffix.lower() != ".xlsx":
        fail("output", "ต้องลงท้าย .xlsx")
    with tempfile.NamedTemporaryFile(prefix=".research-author-", suffix=".xlsx", dir=output.parent, delete=False) as temp:
        temp_path = Path(temp.name)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, output)
    finally:
        temp_path.unlink(missing_ok=True)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--init", action="store_true", help="สร้าง JSON เปล่าสำหรับกรอกข้อมูล")
    group.add_argument("--input", type=Path, help="ไฟล์ JSON ที่ตรวจข้อมูลแล้ว")
    parser.add_argument("--name", help="ชื่อผู้วิจัยเมื่อใช้ --init")
    parser.add_argument("--output", type=Path, help="ไฟล์ปลายทาง .json หรือ .xlsx")
    parser.add_argument("--check", action="store_true", help="ตรวจข้อมูลและจำนวนแถวโดยไม่สร้าง Excel")
    parser.add_argument("--summary-only", action="store_true", help="ใช้กับ --check; แสดงจำนวนคำเตือนเท่านั้น")
    parser.add_argument("--allow-partial", action="store_true", help="ส่งออกรายงาน Scopus ว่างที่ผู้ใช้ยอมรับว่าไม่ครบแล้วเท่านั้น")
    args = parser.parse_args()
    try:
        if args.summary_only and not args.check:
            fail("--summary-only", "ต้องใช้พร้อม --check")
        if args.init:
            if args.check or not args.output:
                fail("--init", "ต้องมี --output และห้ามใช้ --check")
            name = require_text(args.name, "--name")
            if args.output.suffix.lower() != ".json":
                fail("--output", "เมื่อใช้ --init ต้องลงท้าย .json")
            if args.output.exists():
                fail("--output", "ไฟล์มีอยู่แล้ว; เลือกชื่อใหม่เพื่อไม่ให้ข้อมูลเดิมหาย")
            template = {
                "researcher": {"name": name, "affiliation": ""},
                "retrieved_at": date.today().isoformat(),
                "scopus": [],
                "tci": [],
            }
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(json.dumps(template, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            print(f"Created {args.output}")
        else:
            if args.name:
                fail("--name", "ใช้ได้เฉพาะ --init")
            payload = json.loads(args.input.read_text(encoding="utf-8"))
            valid = validate_payload(payload)
            summary = preflight_summary(valid)
            if args.check:
                if args.output:
                    fail("--output", "ห้ามใช้พร้อม --check")
                if args.summary_only:
                    summary.pop("warnings")
                print(json.dumps(summary, ensure_ascii=False, indent=2))
            else:
                if not args.output:
                    fail("--output", "ต้องระบุปลายทาง .xlsx")
                export(valid, args.output, allow_partial=args.allow_partial)
                print(f"Created {args.output}: Scopus {len(valid['scopus'])}, TCI {len(valid['tci'])}; warnings {summary['warning_count']}")
    except (ValueError, OSError, json.JSONDecodeError, RuntimeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())

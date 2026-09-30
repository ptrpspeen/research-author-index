#!/usr/bin/env python3
"""Refresh only changed Scopus EIDs/DOIs and ThaiJO article URLs/DOIs.

Accepts a previous plugin JSON or two-sheet Excel workbook. New CSV/draft files
must cover the complete current result set for each database being refreshed.
Unchanged Q/Tier evidence is reused; changed source/year needs review.
"""

from __future__ import annotations

import argparse
import copy
import json
import sys
from datetime import date, datetime
from pathlib import Path

from export_excel import (COMMON, SCOPUS_EXTRA, TCI_EXTRA, TAIL, URL_FIELDS,
                          annotate_scopus_doi_conflicts, check_date, check_orcid, certified_year_range,
                          is_journal_document, validate_payload)
from import_scopus_csv import load_csv_records
from metric_cache import apply_cache, collect_cache, issn_key
from reconcile_records import doi_key, title_key


Q_FIELDS = ("quartile", "citescore_percentile", "quartile_year", "quartile_subject", "quartile_source_url")
TIER_FIELDS = ("tier", "tier_period", "tier_source_url", "tier_evidence_basis")
IDENTITY_FIELDS = ("doi", "title", "publication_year", "journal", "issn", "volume", "issue", "authors", "document_type")


def excel_value(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return value


def load_excel(path: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("ต้องติดตั้ง openpyxl เพื่ออ่านไฟล์ Excel เก่า") from exc
    workbook = load_workbook(path, read_only=False, data_only=True)
    if workbook.sheetnames != ["Scopus", "TCI"]:
        raise ValueError("Excel เก่าต้องมีสองชีต Scopus และ TCI ตามรูปแบบปลั๊กอิน")
    payload = {"researcher": {}, "retrieved_at": "", "scopus": [], "tci": []}
    for database, columns in (("scopus", COMMON + SCOPUS_EXTRA + TAIL),
                              ("tci", COMMON + TCI_EXTRA + TAIL)):
        sheet = workbook["Scopus" if database == "scopus" else "TCI"]
        headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
        if "Tier วารสาร" in headers and "Tier ในปีตีพิมพ์" not in headers:
            headers["Tier ในปีตีพิมพ์"] = headers["Tier วารสาร"]
        required = {"ชื่อบทความ", "ปีตีพิมพ์ (ค.ศ.)", "ผู้แต่งทั้งหมด", "รายชื่อผู้แต่งครบ",
                    "วารสาร/แหล่งตีพิมพ์", "ลิงก์บทความ", "สถานะตรวจสอบ"}
        if not required <= set(headers):
            raise ValueError(f"ชีต {sheet.title} ไม่ใช่รูปแบบ Excel ของปลั๊กอิน; ขาด {sorted(required-set(headers))}")
        for row_num in range(2, sheet.max_row + 1):
            row = {}
            for field, label, _ in columns:
                if label not in headers:
                    continue
                cell = sheet.cell(row_num, headers[label])
                value = excel_value(cell.hyperlink.target if field in URL_FIELDS and cell.hyperlink else cell.value)
                if value is not None and value != "":
                    row[field] = value
            if not row.get("title"):
                continue
            if not payload["researcher"]:
                payload["researcher"] = {
                    "name": row.get("researcher", ""),
                    "affiliation": row.get("affiliation", ""),
                }
                if row.get("researcher_orcid"):
                    payload["researcher"]["orcid"] = row["researcher_orcid"]
            payload["retrieved_at"] = payload["retrieved_at"] or row.get("retrieved_at", "")
            row.pop("researcher", None)
            row.pop("affiliation", None)
            row.pop("researcher_orcid", None)
            row.pop("researcher_orcid_url", None)
            row.pop("retrieved_at", None)
            row["authors"] = [name.strip() for name in str(row.get("authors", "")).split(";") if name.strip()]
            row["authors_complete"] = row.get("authors_complete", "ครบ") in {"ครบ", "true", True}
            row["publication_year"] = int(row["publication_year"])
            for field in ("author_h_index", "citations", "citescore_percentile", "quartile_year"):
                if field in row:
                    row[field] = int(row[field])
            payload[database].append(row)
    validate_payload(payload)
    return payload


def load_previous(path: Path | None, name: str | None, affiliation: str, today: str) -> tuple[dict, str]:
    if path is None:
        if not name:
            raise ValueError("ไม่มีไฟล์เก่า: ต้องระบุ --name เพื่อสร้างรายงานใหม่")
        return {"researcher": {"name": name, "affiliation": affiliation},
                "retrieved_at": today, "scopus": [], "tci": []}, "new"
    if path.suffix.lower() == ".xlsx":
        payload = load_excel(path)
        if name and title_key(name) != title_key(payload["researcher"]["name"]):
            raise ValueError("ชื่อผู้วิจัยไม่ตรงกับรายงานเก่า; ตรวจตัวตน/ชื่อสะกดก่อนใช้ --previous")
        return payload, "incremental"
    if path.suffix.lower() == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        validate_payload(payload)
        if name and title_key(name) != title_key(payload["researcher"]["name"]):
            raise ValueError("ชื่อผู้วิจัยไม่ตรงกับรายงานเก่า; ตรวจตัวตน/ชื่อสะกดก่อนใช้ --previous")
        return payload, "incremental"
    raise ValueError("ไฟล์เก่าต้องเป็น .json หรือ .xlsx จากปลั๊กอินนี้")


def comparable(field: str, value):
    if field == "doi":
        return doi_key(value or "")
    if field == "title":
        return title_key(value or "")
    if field == "journal":
        return title_key(value or "")
    if field == "issn":
        return issn_key(value or "")
    if field == "authors":
        return [title_key(name) for name in value or []]
    if field == "document_type":
        return (value or "").casefold()
    return value


def changed_fields(old: dict, fresh: dict) -> list[str]:
    return [field for field in IDENTITY_FIELDS
            if comparable(field, old.get(field)) != comparable(field, fresh.get(field))]


def source_same(old: dict, fresh: dict) -> bool:
    return (issn_key(old.get("issn", "")) == issn_key(fresh.get("issn", ""))
            and title_key(old.get("journal", "")) == title_key(fresh.get("journal", ""))
            and old["publication_year"] == fresh["publication_year"])


def append_note(original: str, addition: str) -> str:
    return "; ".join(item for item in (original.strip(), addition) if item)


def refresh_scopus(previous: list[dict], fresh: list[dict], accept_removed: bool) -> tuple[list[dict], dict]:
    old_by_eid = {row.get("scopus_eid"): row for row in previous}
    old_doi: dict[str, list[dict]] = {}
    for row in previous:
        if row.get("doi"):
            old_doi.setdefault(doi_key(row["doi"]), []).append(row)
    if len(old_by_eid) != len(previous) or None in old_by_eid:
        raise ValueError("รายงานเก่ามี Scopus EID ว่างหรือซ้ำ; ตรวจไฟล์เก่าก่อน")
    fresh_eids = {row["scopus_eid"] for row in fresh}
    fresh_dois: dict[str, list[dict]] = {}
    for row in fresh:
        if row.get("doi"):
            fresh_dois.setdefault(doi_key(row["doi"]), []).append(row)

    doi_fallbacks: dict[str, dict] = {}
    fallback_old_eids: set[str] = set()
    for key, old_group in old_doi.items():
        fresh_group = fresh_dois.get(key, [])
        if len(old_group) != 1 or len(fresh_group) != 1:
            continue
        old, new = old_group[0], fresh_group[0]
        old_record_eid, new_eid = old["scopus_eid"], new["scopus_eid"]
        # Exact EID matches have priority on both sides. A DOI fallback cannot
        # steal an old row already represented by another fresh EID, nor can it
        # consume a fresh row that already has an exact old-EID match.
        if old_record_eid in fresh_eids or new_eid in old_by_eid:
            continue
        doi_fallbacks[new_eid] = old
        fallback_old_eids.add(old_record_eid)

    missing = sorted(eid for eid, old in old_by_eid.items()
                     if eid not in fresh_eids and eid not in fallback_old_eids)
    if missing and not accept_removed:
        raise ValueError("CSV ใหม่ไม่มี EID เก่า: " + ", ".join(missing) +
                         "; ตรวจตัวกรอง/โปรไฟล์ก่อน หากยืนยันให้ใช้ --accept-removed")
    result, added, changed, unchanged = [], [], [], 0
    for row in fresh:
        eid = row["scopus_eid"]
        old = old_by_eid.get(eid)
        reindexed = old is None and eid in doi_fallbacks
        if reindexed:
            old = doi_fallbacks[eid]
        if old is None:
            key = doi_key(row["doi"]) if row.get("doi") else ""
            if key and old_doi.get(key):
                row["verification_status"] = "needs_review"
                row["notes"] = append_note(
                    row.get("notes", ""),
                    "Ambiguous DOI update match; previous metrics were not transferred; review record identity.",
                )
            added.append({"eid": eid, "doi": row.get("doi", ""),
                          "metric_lookup": is_journal_document(row)})
            result.append(row)
            continue
        differences = changed_fields(old, row)
        if reindexed:
            differences.insert(0, "scopus_eid")
        if source_same(old, row):
            for field in Q_FIELDS:
                if field in old:
                    row[field] = old[field]
        if row.get("author_h_index") is None and old.get("author_h_index") is not None:
            row["author_h_index"] = old["author_h_index"]
            row["h_index_as_of"] = old["h_index_as_of"]
        if differences:
            row["verification_status"] = "needs_review"
            old_note = old.get("notes", "") if source_same(old, row) else ""
            row["notes"] = append_note(old_note, "ข้อมูลเปลี่ยน: " + ", ".join(differences))
            changed.append({"eid": eid, "doi": row.get("doi", ""), "changed_fields": differences,
                            "metric_lookup": is_journal_document(row)
                            and not bool(row.get("quartile"))})
        else:
            row["verification_status"] = old.get("verification_status", "confirmed")
            if old.get("notes"):
                row["notes"] = old["notes"]
            unchanged += 1
        result.append(row)
    annotate_scopus_doi_conflicts(result)
    return result, {"added": added, "changed": changed, "removed": missing, "unchanged_count": unchanged}


def refresh_tci(previous: list[dict], draft_rows: list[dict], accept_removed: bool) -> tuple[list[dict], dict]:
    old_url = {row["article_url"].rstrip("/"): row for row in previous}
    old_doi = {doi_key(row["doi"]): row for row in previous if row.get("doi")}
    if len(old_url) != len(previous):
        raise ValueError("รายงานเก่ามี ThaiJO URL ซ้ำ")
    fresh_urls = {row["article_url"].rstrip("/") for row in draft_rows}
    fresh_dois = {doi_key(row["doi"]) for row in draft_rows if row.get("doi")}
    missing = sorted(url for url, old in old_url.items()
                     if url not in fresh_urls and (not old.get("doi") or doi_key(old["doi"]) not in fresh_dois))
    if missing and not accept_removed:
        raise ValueError(f"ThaiJO draft ใหม่ไม่มี URL เก่า {len(missing)} รายการ; ตรวจว่ารวบรวมครบก่อน หรือใช้ --accept-removed")
    result, added, changed, unchanged = [], [], [], 0
    for draft in draft_rows:
        row = copy.deepcopy(draft)
        # ThaiJO author profiles may include editorials and other contributions.
        # A citation record alone does not establish that this is a research article.
        row.setdefault("document_type", "")
        row.setdefault("verification_status", "needs_review")
        old = old_url.get(row["article_url"].rstrip("/"))
        if old is None and row.get("doi"):
            old = old_doi.get(doi_key(row["doi"]))
        if old is None:
            added.append({"url": row["article_url"], "doi": row.get("doi", ""), "metric_lookup": True})
            result.append(row)
            continue
        differences = changed_fields(old, row)
        if old["article_url"].rstrip("/") != row["article_url"].rstrip("/"):
            differences.insert(0, "article_url")
        same_journal = (issn_key(old.get("issn", "")) == issn_key(row.get("issn", ""))
                        and title_key(old["journal"]) == title_key(row["journal"]))
        if same_journal:
            for field in ("current_tier", "current_tier_checked_at", "current_tier_source_url"):
                if old.get(field): row[field] = old[field]
        years = certified_year_range(old.get("tier_period", ""))
        if same_journal and years and years[0] <= row["publication_year"] <= years[1]:
            for field in TIER_FIELDS:
                if field in old:
                    row[field] = old[field]
        if differences:
            row["verification_status"] = "needs_review"
            row["notes"] = append_note(old.get("notes", "") if same_journal else "",
                                        "ข้อมูลเปลี่ยน: " + ", ".join(differences))
            changed.append({"url": row["article_url"], "doi": row.get("doi", ""),
                            "changed_fields": differences, "metric_lookup": not bool(row.get("tier"))})
        else:
            row["verification_status"] = old.get("verification_status", "confirmed")
            if old.get("notes"):
                row["notes"] = old["notes"]
            unchanged += 1
        result.append(row)
    return result, {"added": added, "changed": changed, "removed": missing, "unchanged_count": unchanged}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--previous", type=Path, help="ไฟล์เก่า JSON หรือ XLSX; ถ้าไม่มีให้สร้างใหม่")
    parser.add_argument("--name", help="ชื่อผู้วิจัย เมื่อไม่มีไฟล์เก่า")
    parser.add_argument("--affiliation", default="")
    parser.add_argument("--scopus-csv", type=Path, action="append", help="Scopus CSV ครบชุด; ระบุซ้ำได้")
    parser.add_argument("--scopus-filters", default="ไม่บันทึก; ยังไม่ยืนยันตัวกรอง", help="ตัวกรองที่สังเกตจริงบน Documents")
    parser.add_argument("--tci-draft", type=Path, help="ผล collect_thaijo.py จาก URL ครบชุด")
    parser.add_argument("--tci-query", help="คำค้นที่ใช้จริงใน ThaiJO Author Search เพื่อบันทึก audit")
    parser.add_argument("--tci-pages-checked", help="จำนวนหน้าหรือช่วงหน้าที่ตรวจจริง")
    parser.add_argument("--author-id", help="Scopus Author ID ที่ยืนยันแล้ว")
    parser.add_argument("--profile-url")
    parser.add_argument("--h-index", type=int)
    parser.add_argument("--orcid", help="ORCID ที่ยืนยันแล้ว")
    parser.add_argument("--expected-count", type=int, help="จำนวน Scopus Documents ตามโปรไฟล์/ตัวกรอง")
    parser.add_argument("--tci-expected-count", type=int, help="จำนวน ThaiJO article URL ที่ตั้งใจเก็บ")
    parser.add_argument("--retrieved-at", default=date.today().isoformat())
    parser.add_argument("--accept-removed", action="store_true", help="ใช้เมื่อยืนยันว่ารายการเก่าหายจากฐานจริง")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--changes", type=Path, required=True)
    args = parser.parse_args()
    try:
        check_date(args.retrieved_at, "--retrieved-at")
        if not args.scopus_csv and not args.tci_draft:
            raise ValueError("ต้องมี --scopus-csv หรือ --tci-draft อย่างน้อยหนึ่งอย่าง")
        if args.h_index is not None and args.h_index < 0:
            raise ValueError("--h-index ต้องไม่ติดลบ")
        for field in ("expected_count", "tci_expected_count"):
            value = getattr(args, field)
            if value is not None and value < 0:
                raise ValueError(f"--{field.replace('_','-')} ต้องไม่ติดลบ")
        if args.previous and args.output.resolve() == args.previous.resolve():
            raise ValueError("--output ต้องเป็นไฟล์ใหม่ เพื่อเก็บไฟล์เก่าไว้ตรวจทาน")
        if args.output.resolve() == args.changes.resolve():
            raise ValueError("--output และ --changes ต้องเป็นคนละไฟล์")
        payload, mode = load_previous(args.previous, args.name, args.affiliation, args.retrieved_at)
        old_cache = collect_cache(validate_payload(payload), strict=False) if mode == "incremental" else {"scopus": [], "tci": []}
        if args.orcid:
            payload["researcher"]["orcid"] = check_orcid(args.orcid, "--orcid")
        changes = {"mode": mode, "scopus": {"source": "not_refreshed"},
                   "tci": {"source": "not_refreshed"}}
        audit = payload.setdefault("audit", {})
        searches = audit.setdefault("searches", [])
        if args.scopus_csv:
            if args.expected_count is None:
                raise ValueError("Scopus CSV ต้องมี --expected-count จากจำนวนที่สังเกตใน UI")
            previous_ids = {row.get("scopus_author_id") for row in payload["scopus"] if row.get("scopus_author_id")}
            author_id = args.author_id or (next(iter(previous_ids)) if len(previous_ids) == 1 else None)
            if not author_id or not author_id.isdigit():
                raise ValueError("ต้องระบุ --author-id ที่ยืนยันจาก Scopus Author Profile")
            if previous_ids and previous_ids != {author_id}:
                raise ValueError("Author ID ใหม่ไม่ตรงกับไฟล์เก่า; ตรวจว่าเป็นผู้วิจัยคนเดียวกัน")
            profile_url = args.profile_url or f"https://www.scopus.com/authid/detail.uri?authorId={author_id}"
            fresh = load_csv_records(args.scopus_csv, author_id, profile_url, args.h_index,
                                     args.retrieved_at, args.expected_count)
            payload["scopus"], changes["scopus"] = refresh_scopus(payload["scopus"], fresh, args.accept_removed)
            changes["scopus"]["source"] = "csv"
            audit["scopus_expected_count"] = args.expected_count
            audit["scopus_excluded_count"] = 0
            entry = {"database": "scopus", "query": f"AU-ID({author_id})",
                     "method": "Documents CSV export", "results": len(fresh),
                     "pages_checked": "CSV เทียบจำนวนกับ UI", "filters": args.scopus_filters, "checked_at": args.retrieved_at}
            if entry not in searches:
                searches.append(entry)
        if args.tci_draft:
            draft = json.loads(args.tci_draft.read_text(encoding="utf-8"))
            if draft.get("errors"):
                raise ValueError("ThaiJO draft มี fetch errors; แก้ก่อนอัปเดต")
            rows = draft.get("rows")
            if not isinstance(rows, list) or (not rows and args.tci_expected_count != 0):
                raise ValueError("ThaiJO draft ไม่มี rows")
            if args.tci_expected_count is not None and len(rows) != args.tci_expected_count:
                raise ValueError("จำนวน ThaiJO draft ไม่ตรงกับ --tci-expected-count")
            payload["tci"], changes["tci"] = refresh_tci(payload["tci"], rows, args.accept_removed)
            changes["tci"]["source"] = "thaijo_metadata"
            if args.tci_expected_count is not None:
                audit["tci_expected_count"] = args.tci_expected_count
            else:
                audit.pop("tci_expected_count", None)
            audit["tci_excluded_count"] = 0
            if args.tci_query:
                entry = {"database": "tci", "query": args.tci_query,
                         "method": "Author profile + ThaiJO citation metadata", "results": len(rows),
                         "pages_checked": args.tci_pages_checked or "ไม่บันทึก; ยังไม่ยืนยันความครบ",
                         "filters": "ไม่บันทึก", "checked_at": args.retrieved_at}
                if entry not in searches:
                    searches.append(entry)
        payload["retrieved_at"] = args.retrieved_at
        payload, metrics_reused = apply_cache(payload, old_cache)
        changes["summary"] = {
            "researcher": payload["researcher"]["name"],
            "scopus_rows": len(payload["scopus"]), "tci_rows": len(payload["tci"]),
            "scopus_added": len(changes["scopus"].get("added", [])),
            "scopus_changed": len(changes["scopus"].get("changed", [])),
            "tci_added": len(changes["tci"].get("added", [])),
            "tci_changed": len(changes["tci"].get("changed", [])),
            "review_items": len(changes["scopus"].get("changed", [])) + len(changes["tci"].get("changed", []))
                            + len(changes["tci"].get("added", [])),
            "metrics_reused": metrics_reused,
        }
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.changes.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        args.changes.write_text(json.dumps(changes, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(changes["summary"], ensure_ascii=False))
    except (ValueError, OSError, json.JSONDecodeError, KeyError, TypeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())

"""Offline correctness gates; no browser, credentials or network required."""
import copy
import csv
import json
import os
import subprocess
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

ROOT = Path(os.environ.get('RAI_PLUGIN_ROOT', Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(ROOT / 'scripts'))
from export_excel import check_orcid, export, preflight_summary, validate_payload, validate_record
from import_scopus_csv import (REQUIRED, csv_record, load_csv_records, merge_previous,
                               normalize_url, parse_authors)
from incremental_update import load_excel, load_previous, refresh_tci, refresh_scopus
from metric_cache import apply_cache, collect_cache, queue
from verify_thaijo_links import check_one
import collect_thaijo


def scopus(**extra):
    row = dict(title='Example study', publication_year=2021, document_type='Article',
        authors=['Example, Alice'], authors_complete=True, journal='Example Journal', issn='12345679',
        doi='10.0000/example', scopus_eid='2-s2.0-12345', scopus_author_id='123',
        author_profile_url='https://www.scopus.com/authid/detail.uri?authorId=123',
        article_url='https://www.scopus.com/pages/publications/12345', verification_status='confirmed')
    row.update(extra)
    return row


def tci(**extra):
    row = dict(title='Example study', publication_year=2026, document_type='Article',
        authors=['Example Alice'], authors_complete=True, journal='Example Journal', issn='12345679',
        article_url='https://so01.tci-thaijo.org/index.php/example/article/view/1', verification_status='confirmed')
    row.update(extra)
    return row


def metric(**extra):
    value = dict(quartile='Q2', citescore_percentile=60, quartile_year=2021,
        quartile_subject='Computer Science', quartile_source_url='https://www.scopus.com/sourceid/1')
    value.update(extra)
    return value


def tier():
    return dict(tier='Tier 1', tier_period='2025-01-01 – 2029-12-31',
        tier_source_url='https://www.tci-thaijo.org/public/tinfo.php?path=example')


def payload(s=None, t=None):
    return dict(researcher={'name':'Example Alice'}, retrieved_at=date.today().isoformat(),
        scopus=s or [], tci=t or [])


class Correctness(unittest.TestCase):
    def test_reject_other_author_csv(self):
        with self.assertRaises(ValueError): parse_authors('Example, Alice (456)', '456', '123')

    def test_unknown_citations_stay_unknown(self):
        raw={'EID':'2-s2.0-12345','Year':'2021','Cited by':'','Document Type':'Article',
            'Publication Stage':'Final','Title':'Example','Author full names':'Example, Alice (123)',
            'Author(s) ID':'123','Source title':'Journal','ISSN':'12345679','DOI':'',
            'Link':'https://www.scopus.com/pages/publications/12345'}
        self.assertIsNone(csv_record(raw,'123','https://www.scopus.com/authid/detail.uri?authorId=123',None,'2026-09-29').get('citations'))
        raw['Cited by']='0'
        self.assertEqual(csv_record(raw,'123','https://www.scopus.com/authid/detail.uri?authorId=123',None,'2026-09-29')['citations'],0)

    def test_reject_eid_substring_collision(self):
        with self.assertRaises(ValueError): validate_record(scopus(article_url='https://www.scopus.com/pages/publications/123456'),'scopus',0)

    def test_reject_scopus_search_as_article(self):
        with self.assertRaises(ValueError): validate_record(scopus(article_url='https://www.scopus.com/results?q=12345'),'scopus',0)

    def test_reject_profile_substring_collision(self):
        with self.assertRaises(ValueError): validate_record(scopus(author_profile_url='https://www.scopus.com/authid/detail.uri?authorId=1234'),'scopus',0)

    def test_legacy_export_link_supported_exactly(self):
        self.assertEqual(normalize_url('https://www.scopus.com/record/display.uri?eid=2-s2.0-12345&origin=resultslist','2-s2.0-12345'),'https://www.scopus.com/pages/publications/12345')

    def test_current_inward_csv_link_accepts_only_matching_eid(self):
        url = 'https://www.scopus.com/inward/record.uri?eid=2-s2.0-12345&partnerID=40'
        self.assertEqual(normalize_url(url, '2-s2.0-12345'),
                         'https://www.scopus.com/pages/publications/12345')
        for invalid in (url.replace('12345', '123456'), url + '&eid=2-s2.0-999',
                        url.replace('www.scopus.com', 'example.com'),
                        url.replace('/inward/record.uri', '/inward/search.uri')):
            with self.subTest(url=invalid), self.assertRaises(ValueError):
                normalize_url(invalid, '2-s2.0-12345')

    def test_metric_summary_does_not_count_conferences_as_missing_q(self):
        rows = [scopus(**metric()),
                scopus(scopus_eid='2-s2.0-20001', article_url='https://www.scopus.com/pages/publications/20001', doi='10.0000/a'),
                scopus(scopus_eid='2-s2.0-20002', article_url='https://www.scopus.com/pages/publications/20002',
                       doi='10.0000/b', document_type='Conference Paper')]
        summary = preflight_summary(validate_payload(payload(rows)))
        self.assertEqual(summary['scopus_q_eligible_rows'], 2)
        self.assertEqual(summary['scopus_q_evidenced_rows'], 1)
        self.assertEqual(summary['scopus_q_missing_eligible_rows'], 1)
        self.assertEqual(summary['scopus_q_not_applicable_rows'], 1)

    def test_orcid_url_normalization_preserves_checksum_validation(self):
        identifier = '0000-0003-3677-8327'
        for value in (identifier, 'https://orcid.org/' + identifier,
                      'https://orcid.org/' + identifier + '/'):
            self.assertEqual(check_orcid(value, 'test'), identifier)
        for invalid in ('https://orcid.org/' + identifier[:-1] + '8',
                        'https://orcid.org.example.com/' + identifier,
                        'https://user@orcid.org/' + identifier,
                        'https://orcid.org/' + identifier + '?other=1',
                        '0000-0003-3677-X327'):
            with self.subTest(value=invalid), self.assertRaises(ValueError):
                check_orcid(invalid, 'test')

    def test_old_paper_cannot_take_previous_year_q(self):
        with self.assertRaises(ValueError): validate_payload(payload([scopus(**metric(quartile_year=2020))]))

    def test_current_year_previous_q_requires_note(self):
        y=date.today().year
        with self.assertRaises(ValueError): validate_payload(payload([scopus(publication_year=y, **metric(quartile_year=y-1))]))
        validate_payload(payload([scopus(publication_year=y, notes='Latest completed CiteScore year verified in Sources.', **metric(quartile_year=y-1))]))

    def test_snapshot_tier_is_not_historical(self):
        with self.assertRaises(ValueError): validate_payload(payload(t=[tci(publication_year=2021,**{**tier(),'tier_period':'ตรวจ ณ 2026-09-29'})]))

    def test_tier_cache_skips_outside_period(self):
        cache=collect_cache(validate_payload(payload(t=[tci(**tier())])))
        result, filled=apply_cache(payload(t=[tci(publication_year=2021)]),cache)
        self.assertFalse(result['tci'][0].get('tier'))
        self.assertEqual(filled['tci'],0)

    def test_changed_year_drops_invalid_tier(self):
        result,_=refresh_tci([tci(**tier())],[tci(publication_year=2021)],False)
        self.assertFalse(result[0].get('tier'))

    def test_q_cache_needs_matching_subject(self):
        cache=collect_cache(validate_payload(payload([scopus(**metric())])))
        result,filled=apply_cache(payload([scopus()]),cache)
        self.assertFalse(result['scopus'][0].get('quartile'))
        result,filled=apply_cache(payload([scopus(quartile_subject='Computer Science')]),cache)
        self.assertEqual(result['scopus'][0]['quartile'],'Q2')

    def test_q_cache_preserves_multiple_subjects(self):
        a=scopus(**metric())
        b=scopus(title='Another study',doi='10.0000/another',scopus_eid='2-s2.0-6789',
          article_url='https://www.scopus.com/pages/publications/6789',**metric(quartile='Q1',citescore_percentile=80,quartile_subject='Mathematics'))
        cache=collect_cache(validate_payload(payload([a,b])))
        self.assertEqual(len(cache['scopus']),2)

    def test_author_names_must_match_not_just_count(self):
        self.assertTrue(check_one(tci(),tci(authors=['Wrong Person'])))

    def test_review_documents_get_q_queue(self):
        self.assertEqual(len(queue(validate_payload(payload([scopus(document_type='Review')])) )['scopus_issn_year_to_check']),1)

    def test_removed_records_stop(self):
        with self.assertRaises(ValueError):refresh_scopus([scopus()],[],False)

    def test_fresh_duplicate_doi_does_not_transfer_old_metrics(self):
        doi = '10.0000/shared'
        old = scopus(doi=doi, **metric())
        fresh = [
            scopus(title='First record', doi=doi, scopus_eid='2-s2.0-20001',
                   article_url='https://www.scopus.com/pages/publications/20001'),
            scopus(title='Second record', doi=doi, scopus_eid='2-s2.0-20002',
                   article_url='https://www.scopus.com/pages/publications/20002'),
        ]
        rows, changes = refresh_scopus([old], fresh, True)
        self.assertEqual(len(rows), 2)
        self.assertEqual(changes['removed'], [old['scopus_eid']])
        self.assertEqual(len(changes['added']), 2)
        for row in rows:
            self.assertNotIn('quartile', row)
            self.assertEqual(row['verification_status'], 'needs_review')
            self.assertIn('unique works', row['notes'])

    def test_previous_duplicate_doi_does_not_transfer_old_metrics(self):
        doi = '10.0000/shared'
        previous = [
            scopus(doi=doi, scopus_eid='2-s2.0-10001',
                   article_url='https://www.scopus.com/pages/publications/10001', **metric()),
            scopus(title='Second old record', doi=doi, scopus_eid='2-s2.0-10002',
                   article_url='https://www.scopus.com/pages/publications/10002', **metric(quartile='Q1')),
        ]
        fresh = [scopus(title='Reindexed record', doi=doi, scopus_eid='2-s2.0-20001',
                        article_url='https://www.scopus.com/pages/publications/20001')]
        rows, changes = refresh_scopus(previous, fresh, True)
        self.assertEqual(changes['removed'], ['2-s2.0-10001', '2-s2.0-10002'])
        self.assertEqual(len(changes['added']), 1)
        self.assertNotIn('quartile', rows[0])
        self.assertEqual(rows[0]['verification_status'], 'needs_review')
        self.assertIn('Ambiguous DOI update match', rows[0]['notes'])

    def test_exact_eid_matches_take_priority_over_doi_fallback(self):
        doi = '10.0000/shared'
        old = scopus(doi=doi, **metric())
        exact = scopus(doi='10.0000/changed', **metric(quartile='Q1'))
        fresh_other = scopus(title='Other current record', doi=doi, scopus_eid='2-s2.0-6789',
                             article_url='https://www.scopus.com/pages/publications/6789')
        rows, _ = refresh_scopus([old], [exact, fresh_other], False)
        self.assertEqual(rows[0]['scopus_eid'], old['scopus_eid'])
        self.assertEqual(rows[0]['quartile'], 'Q2')
        self.assertNotIn('quartile', rows[1])
        self.assertEqual(rows[1]['verification_status'], 'needs_review')

        old_other = scopus(title='Other old record', doi='10.0000/other',
                           scopus_eid='2-s2.0-6789', article_url='https://www.scopus.com/pages/publications/6789',
                           **metric(quartile='Q1'))
        current = scopus(title='Other old record', doi=doi,
                         scopus_eid='2-s2.0-6789', article_url='https://www.scopus.com/pages/publications/6789')
        rows, changes = refresh_scopus([old, old_other], [current], True)
        self.assertEqual(changes['removed'], [old['scopus_eid']])
        self.assertEqual(rows[0]['quartile'], 'Q1')

    def test_count_mismatch_stops(self):
        with self.assertRaises(ValueError):validate_payload({**payload([scopus()]),'audit':{'scopus_expected_count':2}})

    def test_wrong_previous_researcher_stops(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'previous.json';p.write_text(json.dumps(payload([scopus()])))
            with self.assertRaises(ValueError):load_previous(p,'Different Person','',date.today().isoformat())

    def test_cross_database_overlap_is_preserved(self):
        result=validate_payload(payload([scopus()], [tci()]))
        self.assertEqual((len(result['scopus']),len(result['tci'])),(1,1))

    def test_conflicting_cache_does_not_replace_q(self):
        cache=collect_cache(validate_payload(payload([scopus(**metric())])))
        with self.assertRaises(ValueError):apply_cache(payload([scopus(**metric(quartile='Q1',citescore_percentile=80))]),cache)

    def test_conference_does_not_get_journal_q(self):
        with self.assertRaises(ValueError):validate_payload(payload([scopus(document_type='Conference Paper',**metric())]))

    def test_multiple_tier_periods_are_kept(self):
        old=tci(publication_year=2021, title='Old article', article_url='https://so01.tci-thaijo.org/index.php/example/article/view/2',
          **{**tier(),'tier':'Tier 2','tier_period':'2020 – 2024'})
        cache=collect_cache(validate_payload(payload(t=[tci(**tier()), old])))
        self.assertEqual(len(cache['tci']),2)
        result,_=apply_cache(payload(t=[tci(publication_year=2021)]),cache)
        self.assertEqual(result['tci'][0]['tier'],'Tier 2')

    def test_zero_csv_requires_observed_zero(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'zero.csv'
            with p.open('w',newline='') as f:csv.writer(f).writerow(sorted(REQUIRED))
            self.assertEqual(load_csv_records([p],'123','https://www.scopus.com/authid/detail.uri?authorId=123',None,'2026-09-29',0),[])
            with self.assertRaises(ValueError):load_csv_records([p],'123','',None,'2026-09-29',1)

    def test_counts_are_not_invented_in_tci_import(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d);(p/'draft.json').write_text(json.dumps({'rows':[tci()], 'errors':[]}))
            result=subprocess.run([sys.executable,str(ROOT/'scripts/incremental_update.py'),
                '--name','Example Alice','--tci-draft',str(p/'draft.json'),'--output',str(p/'out.json'),
                '--changes',str(p/'changes.json')],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            audit=json.loads((p/'out.json').read_text())['audit']
            self.assertNotIn('tci_expected_count',audit)

    def test_scopus_import_requires_observed_count(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)
            result=subprocess.run([sys.executable,str(ROOT/'scripts/incremental_update.py'),
                '--name','Example Alice','--scopus-csv',str(p/'unopened.csv'),'--author-id','123',
                '--output',str(p/'out.json'),'--changes',str(p/'changes.json')],capture_output=True,text=True)
            self.assertNotEqual(result.returncode,0)
            self.assertIn('--expected-count',result.stderr)
            self.assertFalse((p/'out.json').exists())

    def test_excel_roundtrip_and_hyperlinks(self):
        from openpyxl import load_workbook
        data=payload([scopus(**metric())],[tci(**tier())])
        data['researcher']['orcid']='0000-0002-1825-0097'
        data['tci'][0]['title']='=formula-like title'
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'report.xlsx'
            export(validate_payload(data),p)
            wb=load_workbook(p)
            self.assertEqual(wb.sheetnames,['Scopus','TCI'])
            for sheet in wb:
                self.assertEqual(sheet.max_row,2)
                headers={c.value:c.column for c in sheet[1]}
                self.assertTrue(sheet.cell(2,headers['ลิงก์บทความ']).hyperlink.target.startswith('https://'))
                self.assertNotEqual(sheet.cell(2,headers['ชื่อบทความ']).data_type,'f')
                self.assertEqual(sheet.cell(2,headers['ORCID ผู้วิจัย']).value,data['researcher']['orcid'])
            restored=load_excel(p)
            self.assertEqual(restored['scopus'][0]['quartile'],'Q2')
            self.assertEqual(restored['tci'][0]['tier'],'Tier 1')

    def test_collector_reports_real_field_completeness(self):
        summary=collect_thaijo.collection_summary([tci(volume='1'),tci(publication_year=None)],[])
        self.assertEqual(summary['rows'],2)
        self.assertEqual(summary['missing_fields']['publication_year'],1)
        self.assertEqual(summary['missing_fields']['volume'],1)

    def test_author_whitespace_is_not_a_different_person(self):
        summary=collect_thaijo.review_draft({'rows':[tci(authors=['Thara  Angskun']),tci(authors=['Jitimon Angskun'])]},['Thara Angskun'])
        self.assertEqual(summary['target_name_matches'],1)
        self.assertEqual(len(summary['unmatched_target_names']),1)
        self.assertEqual(summary['unmatched_target_names'][0]['row'],2)

    def test_unknown_thaijo_type_is_not_invented(self):
        draft=tci(title='From Editor');draft.pop('document_type')
        rows,_=refresh_tci([], [draft], False)
        self.assertFalse(rows[0].get('document_type'))

    def test_blocked_scopus_cannot_silently_export_empty(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'report.xlsx'
            with self.assertRaises(ValueError):export(validate_payload(payload(t=[tci()])),p)
            self.assertFalse(p.exists())

    def test_explicit_partial_is_labeled(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'partial.xlsx'
            export(validate_payload(payload(t=[tci()])),p,allow_partial=True)
            self.assertIn('Scopus ยังไม่ครบ',load_workbook(p).properties.description)

    def test_verified_zero_can_export_empty_scopus(self):
        data=payload(t=[tci()]);data['audit']={'scopus_expected_count':0,'searches':[
            {'database':'scopus','query':'Observed author profile documents','results':0,'checked_at':'2026-09-29'}]}
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'zero.xlsx';export(validate_payload(data),p);self.assertTrue(p.exists())

    def test_same_title_year_distinct_editorials_are_not_deleted(self):
        a=tci(title='From Editor',document_type='Editorial',issue='1')
        b=tci(title='From Editor',document_type='Editorial',issue='2',
            article_url='https://so01.tci-thaijo.org/index.php/example/article/view/2')
        valid=validate_payload(payload(t=[a,b]))
        self.assertEqual(len(valid['tci']),2)

    def test_same_doi_is_still_rejected(self):
        a=tci(doi='10.0000/same')
        b=tci(title='Different title',doi='https://doi.org/10.0000/SAME',
            article_url='https://so01.tci-thaijo.org/index.php/example/article/view/2')
        with self.assertRaises(ValueError):validate_payload(payload(t=[a,b]))

    def test_scopus_duplicate_doi_csv_records_are_retained_and_flagged(self):
        def csv_row(eid, title, year):
            numeric = eid.removeprefix('2-s2.0-')
            return {
                'Author full names': 'Example, Alice (123)', 'Author(s) ID': '123',
                'Title': title, 'Year': str(year), 'Source title': 'Example Journal',
                'Cited by': '0', 'DOI': '10.0000/shared',
                'Link': f'https://www.scopus.com/pages/publications/{numeric}',
                'ISSN': '12345679', 'Document Type': 'Article',
                'Publication Stage': 'Final', 'EID': eid,
            }

        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'duplicate-doi.csv'
            with path.open('w', newline='', encoding='utf-8') as handle:
                writer = csv.DictWriter(handle, fieldnames=sorted(REQUIRED))
                writer.writeheader()
                writer.writerow(csv_row('2-s2.0-20001', 'First record', 2021))
                writer.writerow(csv_row('2-s2.0-20002', 'Second record', 2020))
            rows = load_csv_records([path], '123',
                'https://www.scopus.com/authid/detail.uri?authorId=123', None,
                date.today().isoformat(), 2)

        valid = validate_payload(payload(rows))
        self.assertEqual(len(valid['scopus']), 2)
        self.assertEqual({row['verification_status'] for row in valid['scopus']}, {'needs_review'})
        self.assertTrue(all('unique works' in row['notes'] for row in valid['scopus']))
        summary = preflight_summary(valid)
        self.assertEqual(summary['scopus_doi_conflict_groups'], 1)
        self.assertTrue(any('distinct EIDs' in warning and 'not a unique-work count' in warning
                            for warning in summary['warnings']))
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            output = Path(d) / 'duplicate-doi.xlsx'
            export(valid, output)
            workbook = load_workbook(output)
            self.assertIn('records ไม่ใช่จำนวน unique works', workbook.properties.description)
            headers = {cell.value: cell.column for cell in workbook['Scopus'][1]}
            note_col = headers['หมายเหตุ']
            self.assertTrue(all('unique works' in workbook['Scopus'].cell(row, note_col).value
                                for row in (2, 3)))

    def test_scopus_duplicate_eid_remains_rejected(self):
        a = scopus()
        b = scopus(title='Different title', doi='10.0000/another')
        with self.assertRaises(ValueError):
            validate_payload(payload([a, b]))

    def test_hand_edited_duplicate_doi_is_annotated_and_merge_keeps_note(self):
        rows = [
            scopus(scopus_eid='2-s2.0-20001', article_url='https://www.scopus.com/pages/publications/20001'),
            scopus(title='Second record', scopus_eid='2-s2.0-20002',
                   article_url='https://www.scopus.com/pages/publications/20002'),
        ]
        validated = validate_payload(payload(rows))['scopus']
        self.assertTrue(all(row['verification_status'] == 'needs_review' for row in validated))
        self.assertTrue(all('unique works' in row['notes'] for row in validated))
        previous = [dict(row, notes='') for row in validated]
        merged = merge_previous(validated, previous)
        self.assertTrue(all('unique works' in row['notes'] for row in merged))


if __name__=='__main__':unittest.main()
